"""Mercura AI pitch deck. 16:9, ivory / navy / gold / teal. Facts match the running demo."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

IVORY = RGBColor(0xF5, 0xF0, 0xE7)
NAVY = RGBColor(0x08, 0x1F, 0x2A)
GOLD = RGBColor(0xC9, 0xA4, 0x5C)
TEAL = RGBColor(0x0F, 0x9F, 0x91)
BODY = RGBColor(0x1A, 0x38, 0x44)
CARD = RGBColor(0xFB, 0xF8, 0xF1)
WHITE = RGBColor(0xFF, 0xFC, 0xF7)
RED = RGBColor(0x8E, 0x2F, 0x2F)
W, H = 13.333, 7.5
SERIF, SANS = "Georgia", "Calibri"
IMG = "/Users/anushak/Documents/Projects/Hackathon/Mercura-App-main/Mercura-App-main/img"

prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)
BLANK = prs.slide_layouts[6]
TOTAL = 12


def rgb(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def rect(slide, l, t, w, h, color):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    rgb(s, color)
    s.shadow.inherit = False
    return s


def rrect(slide, l, t, w, h, color, line=None, lw=1):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    rgb(s, color)
    s.adjustments[0] = 0.08
    if line:
        s.line.color.rgb = line
        s.line.width = Pt(lw)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    return s


def text(slide, l, t, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    tf.vertical_anchor = anchor
    for i, (msg, font, size, color, bold) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        r = p.add_run()
        r.text = msg
        r.font.name = font
        r.font.size = Pt(size)
        r.font.color.rgb = color
        r.font.bold = bold
    return box


def bg(slide, color=IVORY):
    rect(slide, 0, 0, W, H, color)


def mark(slide, l, t, size=0.42):
    rrect(slide, l, t, size, size, GOLD)
    text(slide, l, t + size * 0.16, size, size * 0.7, [("M", SERIF, int(size * 36), NAVY, True)], PP_ALIGN.CENTER)


def footer(slide, n):
    rect(slide, 0, 7.16, W, 0.34, NAVY)
    rect(slide, 0, 7.16, W, 0.015, GOLD)
    text(slide, 0.5, 7.2, 2.4, 0.24, [("MERCURA AI", SANS, 10, GOLD, True)], anchor=MSO_ANCHOR.MIDDLE)
    text(slide, 3.2, 7.2, 6.8, 0.24, [("Driven by AI and controlled by you", SANS, 11, IVORY, False)], PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    text(slide, 10.6, 7.2, 2.2, 0.24, [(f"{n:02d}  /  {TOTAL:02d}", SANS, 11, GOLD, True)], PP_ALIGN.RIGHT, MSO_ANCHOR.MIDDLE)


def kicker(slide, n, label):
    text(slide, 0.55, 0.32, 1.1, 0.28, [(f"{n:02d}", SANS, 13, GOLD, True)])
    text(slide, 1.15, 0.32, 8, 0.28, [(label, SANS, 12, NAVY, True)])
    rect(slide, 0.55, 0.68, 1.4, 0.015, GOLD)


def title_block(slide, headline):
    text(slide, 0.55, 0.78, 12.2, 0.55, [(headline, SERIF, 30, NAVY, False)])


# ---------- 1 Title ----------
s = prs.slides.add_slide(BLANK)
bg(s)
rect(s, 0, 0, W, 0.08, NAVY)
rect(s, 0, 0.08, W, 0.012, GOLD)
text(s, 0.6, 0.32, 6, 0.28, [("T.01   FINTECH & COMMERCE", SANS, 12, TEAL, True)])
text(s, 6.4, 0.32, 6.3, 0.28, [("MANIPAL ACADEMY OF HIGHER EDUCATION", SANS, 12, NAVY, True)], PP_ALIGN.RIGHT)
mark(s, 0.6, 1.45, 0.62)
text(s, 1.4, 1.52, 6.6, 0.5, [("MERCURA AI", SERIF, 36, NAVY, False)], anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.6, 2.35, 7.6, 1.15, [("From Dead Stock", SERIF, 46, NAVY, False)])
text(s, 0.6, 3.35, 7.6, 1.0, [("to Real Profit.", SERIF, 46, GOLD, False)])
rect(s, 0.62, 4.5, 2.1, 0.018, GOLD)
text(s, 0.6, 4.7, 7.6, 0.4, [("Smarter Inventory. Healthier Business.", SANS, 18, TEAL, True)])
text(s, 0.6, 5.15, 7.6, 0.35, [("Driven by AI and controlled by you.", SANS, 15, NAVY, False)])
rect(s, 8.55, 1.4, 4.2, 4.35, GOLD)
s.shapes.add_picture(IMG + "/storefront_hero.jpg", Inches(8.62), Inches(1.47), Inches(4.06), Inches(4.21))
rect(s, 0, 6.15, W, 1.35, NAVY)
text(s, 0.6, 6.38, 12, 0.3, [("TEAM STACK OVERFLOW", SANS, 12, GOLD, True)])
text(s, 0.6, 6.72, 12, 0.32, [("Anusha K    ·    Abhishek Anil Kumar    ·    Aryan Singh    ·    Adarsh B Acharya", SANS, 16, IVORY, False)])

# ---------- 2 Problem ----------
s = prs.slides.add_slide(BLANK)
bg(s)
kicker(s, 2, "THE PROBLEM")
title_block(s, "Cash is sitting on the shelf.")
cards = [
    ("01", "Dead stock", "A product can sit for months. The money that bought it cannot buy what customers want now."),
    ("02", "Weak visibility", "A stock count shows what is left. It does not show which items are ageing or slowing."),
    ("03", "Popular is not profit", "100 units at ₹10 profit earn ₹1,000. 40 units at ₹50 profit earn ₹2,000."),
    ("04", "Demand moves", "Festivals, weather, and nearby events change what will sell before the next purchase."),
]
for i, (num, head, body) in enumerate(cards):
    x = 0.55 + i * 3.15
    rrect(s, x, 1.7, 3.0, 3.35, CARD, GOLD, 1.25)
    text(s, x + 0.2, 1.9, 2.6, 0.4, [(num, SERIF, 22, GOLD, False)])
    text(s, x + 0.2, 2.45, 2.6, 0.7, [(head, SERIF, 20, NAVY, False)])
    text(s, x + 0.2, 3.3, 2.6, 1.5, [(body, SANS, 14, BODY, False)])
text(s, 0.55, 5.2, 12.2, 0.28, [("The 100-unit and 40-unit comparison is an illustrative example from the project brief, not a shop result.", SANS, 12, BODY, False)])
rect(s, 0.55, 5.6, 12.23, 1.3, NAVY)
text(s, 0.8, 5.78, 11.7, 0.3, [("IN THE SEEDED DEMO CATALOG", SANS, 11, GOLD, True)])
text(s, 0.8, 6.12, 11.7, 0.5, [("The winter jacket is 63 days old and holds ₹38,000 at cost. That figure is computed from 47 units, not assumed.", SANS, 16, IVORY, False)])
footer(s, 2)

# ---------- 3 Solution ----------
s = prs.slides.add_slide(BLANK)
bg(s)
kicker(s, 3, "OUR SOLUTION")
title_block(s, "Tell the merchant what to do next.")
rrect(s, 0.55, 1.65, 5.7, 5.15, NAVY)
text(s, 0.85, 1.95, 5.1, 1.5, [("Mercura AI turns the shop book into a decision.", SERIF, 26, IVORY, False)])
text(s, 0.85, 3.6, 5.1, 1.6, [("It reads stock, cost, price, age, daily speed, and supplier lead time. It then names what to stop buying, what to discount, what to reorder, and where cash is stuck.", SANS, 16, IVORY, False)])
text(s, 0.85, 5.5, 5.1, 0.9, [("The merchant approves every purchase. Mercura does not buy on its own.", SANS, 15, GOLD, True)])
qs = [
    ("Stop buying", "Slow, ageing stock stays unsold if another order goes out."),
    ("Discount", "A price cut is a choice, with a 30-day cash estimate beside it."),
    ("Reorder", "Fast stock is reordered only when cover is shorter than lead time."),
    ("See the cash", "Capital at risk and a sale-and-purchase ledger stay in view."),
]
for i, (h, b) in enumerate(qs):
    y = 1.65 + i * 1.28
    rrect(s, 6.5, y, 6.25, 1.16, CARD, GOLD, 1)
    rect(s, 6.5, y, 0.08, 1.16, TEAL)
    text(s, 6.8, y + 0.12, 5.7, 0.36, [(h, SERIF, 18, NAVY, False)])
    text(s, 6.8, y + 0.52, 5.7, 0.5, [(b, SANS, 13, BODY, False)])
footer(s, 3)

# ---------- 4 Versus Netstock ----------
s = prs.slides.add_slide(BLANK)
bg(s)
kicker(s, 4, "WHY THIS IS DIFFERENT")
title_block(s, "A shop book, not an ERP add-on.")
text(s, 0.55, 1.42, 12.2, 0.45, [("Netstock plans demand for a business that already runs an ERP. Mercura starts where that shop is not: one owner, little history, and no system of record.", SANS, 15, BODY, False)])
heads = [("Netstock", NAVY, IVORY), ("Mercura AI", TEAL, WHITE)]
for i, (label, fill, ink) in enumerate(heads):
    rrect(s, 3.3 + i * 4.85, 2.0, 4.65, 0.48, fill)
    text(s, 3.3 + i * 4.85, 2.06, 4.65, 0.36, [(label, SANS, 16, ink, True)], PP_ALIGN.CENTER)
rows = [
    ("Starts from", "An ERP such as NetSuite, SAP, or Dynamics", "A sale, a spoken line, or a photo of a bill"),
    ("Who it is for", "Planners across locations and suppliers", "One merchant, in Hindi, Kannada, or English"),
    ("The job", "Forecast, safety stock, and replenishment", "Dead cash, a discount estimate, and a reorder the owner sends"),
    ("Who decides", "Recommendations flow back into the ERP", "Nothing is bought until the owner says so"),
]
for i, (topic, left, right) in enumerate(rows):
    y = 2.6 + i * 0.85
    text(s, 0.55, y + 0.15, 2.55, 0.5, [(topic, SANS, 14, GOLD, True)], anchor=MSO_ANCHOR.MIDDLE)
    rrect(s, 3.3, y, 4.65, 0.75, CARD, GOLD, 1)
    rrect(s, 8.15, y, 4.65, 0.75, CARD, TEAL, 1.25)
    text(s, 3.45, y + 0.12, 4.35, 0.52, [(left, SANS, 13, NAVY, False)], anchor=MSO_ANCHOR.MIDDLE)
    text(s, 8.3, y + 0.12, 4.35, 0.52, [(right, SANS, 13, NAVY, False)], anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.55, 6.1, 12.2, 0.7, [("This is a difference of job. It is not a claim that Mercura forecasts more accurately than Netstock.", SANS, 13, BODY, False)])
footer(s, 4)

# ---------- 5 How it works ----------
s = prs.slides.add_slide(BLANK)
bg(s)
kicker(s, 5, "HOW IT WORKS")
title_block(s, "Collect. Analyse. Recommend. The merchant acts.")
steps = [
    ("1", "Collect", "A sale, a spoken line, or an invoice photo."),
    ("2", "Analyse", "Speed, age, margin, and days of cover."),
    ("3", "Score", "Sold out, dead, risk, low, or healthy."),
    ("4", "Advise", "Discount, wait, or reorder, with the reason."),
    ("5", "Act", "Sell, lock a purchase, or send a supplier note."),
    ("6", "Refresh", "The same figures update across the shop."),
]
rect(s, 0.9, 2.55, 11.5, 0.015, GOLD)
for i, (n, h, b) in enumerate(steps):
    x = 0.55 + i * 2.1
    rrect(s, x + 0.55, 2.28, 0.52, 0.52, TEAL)
    text(s, x + 0.55, 2.34, 0.52, 0.42, [(n, SANS, 16, WHITE, True)], PP_ALIGN.CENTER)
    text(s, x, 3.05, 2.0, 0.45, [(h, SERIF, 18, NAVY, False)], PP_ALIGN.CENTER)
    text(s, x, 3.55, 2.0, 1.3, [(b, SANS, 13, BODY, False)], PP_ALIGN.CENTER)
rrect(s, 0.55, 5.2, 12.23, 1.6, CARD, GOLD, 1)
text(s, 0.85, 5.4, 11.6, 0.3, [("THE LOOP STAYS WITH THE MERCHANT", SANS, 12, GOLD, True)])
text(s, 0.85, 5.8, 11.6, 0.7, [("A recorded sale changes stock, profit, the forecast, and the ledger together. Nothing is ordered until the merchant sends the note or marks stock received.", SANS, 16, NAVY, False)])
footer(s, 5)

# ---------- 6 Features ----------
s = prs.slides.add_slide(BLANK)
bg(s)
kicker(s, 6, "KEY FEATURES")
title_block(s, "Five decisions the shop can make today.")
feats = [
    ("01", "Dead-stock detection", "Age of 60 days or more, and speed under one unit a day, marks stock as dead."),
    ("02", "Smart reorder", "The batch covers supplier lead time plus 14 days of demand, and only when cover is short."),
    ("03", "A reason with the action", "Each row names the product, the money tied up, and whether to discount or reorder."),
    ("04", "Profit beside demand", "The same catalog is sorted by units sold and by rupee profit. They are not the same list."),
    ("05", "Cash and a ledger", "Capital at risk, three discount estimates, and every sale and purchase on one page."),
]
for i, (n, h, b) in enumerate(feats):
    y = 1.55 + i * 0.95
    rrect(s, 0.55, y, 12.23, 0.86, CARD, GOLD, 1)
    text(s, 0.75, y + 0.18, 0.7, 0.5, [(n, SERIF, 18, GOLD, False)], anchor=MSO_ANCHOR.MIDDLE)
    text(s, 1.6, y + 0.1, 4.2, 0.66, [(h, SERIF, 18, NAVY, False)], anchor=MSO_ANCHOR.MIDDLE)
    text(s, 5.9, y + 0.12, 6.6, 0.62, [(b, SANS, 14, BODY, False)], anchor=MSO_ANCHOR.MIDDLE)
footer(s, 6)

# ---------- 7 Rescue ----------
s = prs.slides.add_slide(BLANK)
bg(s)
kicker(s, 7, "MERCURA RESCUE MODE")
title_block(s, "Three prices for the same slow jacket.")
text(s, 0.55, 1.45, 12.2, 0.55, [("Rescue Mode estimates 30-day cash if the winter jacket sells at 10, 15, or 20 percent off. It is an estimate from current speed, not a forecast of what customers will pay.", SANS, 15, BODY, False)])
offers = [
    ("10% off", "₹17,889", "About 13 units"),
    ("15% off", "₹18,845", "About 15 units"),
    ("20% off", "₹19,571", "About 16 units"),
]
for i, (label, money, units) in enumerate(offers):
    x = 0.55 + i * 4.2
    rrect(s, x, 2.2, 3.95, 2.55, NAVY if i == 1 else CARD, GOLD, 1.25)
    ink = GOLD if i == 1 else TEAL
    word = IVORY if i == 1 else NAVY
    sub = IVORY if i == 1 else BODY
    text(s, x + 0.25, 2.4, 3.4, 0.35, [(label, SANS, 14, ink, True)])
    text(s, x + 0.25, 2.85, 3.4, 0.8, [(money, SERIF, 36, word, False)])
    text(s, x + 0.25, 3.85, 3.4, 0.5, [(units, SANS, 16, sub, False)])
text(s, 0.55, 4.95, 12.2, 0.7, [("Demo catalog: Winter Heavy Jacket, 47 units, list price ₹1,499, speed 0.34 a day. A deeper cut moves more units, so the total can rise while each unit earns less.", SANS, 14, BODY, False)])
text(s, 0.55, 5.7, 12.2, 1.05, [("A 15% clearance plan flags that jacket for review. It does not change the shelf price by itself. A bundle price is the next comparison. This build does not calculate bundles.", SANS, 14, NAVY, False)])
footer(s, 7)

# ---------- 8 Product ----------
s = prs.slides.add_slide(BLANK)
bg(s)
kicker(s, 8, "PRODUCT")
title_block(s, "One shop book, on every screen.")
kpis = [("₹1,29,412", "Inventory at cost"), ("₹1,43,428", "Sales on record"), ("₹80,477", "Gross profit"), ("₹38,000", "Dead capital")]
for i, (v, lab) in enumerate(kpis):
    x = 0.55 + i * 3.15
    rrect(s, x, 1.55, 3.0, 1.25, CARD, GOLD, 1)
    text(s, x + 0.15, 1.65, 2.7, 0.55, [(v, SERIF, 22, NAVY, False)], PP_ALIGN.CENTER)
    text(s, x + 0.15, 2.25, 2.7, 0.35, [(lab, SANS, 13, TEAL, True)], PP_ALIGN.CENTER)
rows = [
    ("Winter Heavy Jacket", "JKT-409", "Dead", "63 days", "₹38,000", "jacket.jpg"),
    ("Sports Running Shoes", "SH-102", "Low", "1.7 days cover", "8.4 / day", "shoes.jpg"),
    ("Silk Blend Festive Kurti", "KT-204", "Critical", "Lead time 6 days", "5 on hand", "kurti.jpg"),
    ("Coastal Sandals", "SD-881", "Risk", "Age 44 days", "Soft speed", "sandals.jpg"),
]
rrect(s, 0.55, 3.0, 8.35, 3.75, CARD, GOLD, 1)
text(s, 0.75, 3.15, 7.9, 0.3, [("INVENTORY HEALTH  ·  SEEDED CATALOG", SANS, 11, GOLD, True)])
for i, (name, sku, status, note, extra, photo) in enumerate(rows):
    y = 3.55 + i * 0.75
    s.shapes.add_picture(IMG + "/" + photo, Inches(0.75), Inches(y), Inches(0.55), Inches(0.55))
    text(s, 1.45, y, 2.6, 0.55, [(name, SANS, 13, NAVY, True), (sku, SANS, 11, BODY, False)])
    tone = RED if status in ("Dead", "Critical") else (GOLD if status == "Risk" else TEAL)
    text(s, 4.15, y + 0.05, 1.4, 0.4, [(status, SANS, 13, tone, True)])
    text(s, 5.55, y + 0.05, 3.1, 0.4, [(note + "  ·  " + extra, SANS, 12, BODY, False)])
rrect(s, 9.1, 3.0, 3.68, 3.75, NAVY)
text(s, 9.3, 3.2, 3.3, 0.3, [("SCREENS", SANS, 12, GOLD, True)])
for i, line in enumerate(["Home and vision", "Dashboard", "Inventory", "Sales and forecast", "Alerts and festivals", "Insights and actions", "Cash, orders, ledger", "Copilot  ·  Hindi, Kannada, English"]):
    text(s, 9.3, 3.6 + i * 0.36, 3.3, 0.34, [(line, SANS, 13, IVORY, False)])
footer(s, 8)

# ---------- 9 Architecture ----------
s = prs.slides.add_slide(BLANK)
bg(s)
kicker(s, 9, "ARCHITECTURE")
title_block(s, "What this build actually runs.")
layers = [
    ("See", "One HTML application styled with Tailwind. Home, dashboard, inventory, sales, alerts, insights, cash, and copilot."),
    ("Think", "engine.js scores velocity, age, cover, and status. live.js adds the 14-day pace, reorder quantity, festival lift, and rescue prices."),
    ("Remember", "The seed is data/catalog.json. Sales, new products, clearance, and locked orders stay in this browser."),
    ("Listen", "The copilot answers from the live catalog. Hindi and Kannada are spoken aloud. An invoice photo is read when the words match a product."),
]
for i, (h, b) in enumerate(layers):
    y = 1.55 + i * 1.05
    rrect(s, 0.55, y, 1.55, 0.9, TEAL)
    text(s, 0.55, y + 0.22, 1.55, 0.45, [(h, SANS, 16, WHITE, True)], PP_ALIGN.CENTER)
    rrect(s, 2.25, y, 10.53, 0.9, CARD, GOLD, 1)
    text(s, 2.5, y + 0.18, 10.05, 0.55, [(b, SANS, 14, NAVY, False)], anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.55, 5.85, 12.2, 0.95, [("A Python program only serves the files. React, FastAPI, PostgreSQL, and scikit-learn are the proposed next stack. The risk number in the product is a fixed formula, not a trained model.", SANS, 14, BODY, False)])
footer(s, 9)

# ---------- 10 Impact ----------
s = prs.slides.add_slide(BLANK)
bg(s)
kicker(s, 10, "BUSINESS IMPACT")
title_block(s, "Measured on the demo catalog only.")
stats = [
    ("₹1,29,412", "On-hand inventory", "at cost, 11 products"),
    ("₹1,43,428", "Sales value", "already on record"),
    ("₹80,477", "Gross profit", "on those recorded sales"),
    ("₹38,000", "Dead capital", "one jacket, 63 days"),
]
for i, (v, a, b) in enumerate(stats):
    x = 0.55 + (i % 4) * 3.15
    rrect(s, x, 1.6, 3.0, 2.15, CARD, GOLD, 1.25)
    text(s, x + 0.12, 1.75, 2.76, 0.7, [(v, SERIF, 24, NAVY if i < 3 else RED, False)], PP_ALIGN.CENTER)
    text(s, x + 0.12, 2.5, 2.76, 0.4, [(a, SANS, 14, TEAL, True)], PP_ALIGN.CENTER)
    text(s, x + 0.12, 2.95, 2.76, 0.45, [(b, SANS, 13, BODY, False)], PP_ALIGN.CENTER)
text(s, 0.55, 4.0, 12.2, 0.35, [("SHARE OF INVENTORY VALUE", SANS, 12, GOLD, True)])
bars = [("Healthy stock", 71110, TEAL), ("Ageing, still watched", 20302, GOLD), ("Dead stock", 38000, RED)]
for i, (lab, val, col) in enumerate(bars):
    y = 4.45 + i * 0.55
    text(s, 0.55, y, 2.6, 0.4, [(lab, SANS, 14, NAVY, False)], anchor=MSO_ANCHOR.MIDDLE)
    rect(s, 3.3, y + 0.08, 7.4, 0.26, RGBColor(0xE7, 0xE0, 0xD4))
    rect(s, 3.3, y + 0.08, 7.4 * val / 129412, 0.26, col)
    text(s, 10.85, y, 1.9, 0.4, [(f"₹{val:,}".replace(",", ","), SANS, 14, NAVY, True)], PP_ALIGN.RIGHT, MSO_ANCHOR.MIDDLE)
text(s, 0.55, 6.2, 12.2, 0.7, [("Sharma Apparels seeded catalog. Not a customer count, not a field saving, and not a performance claim. Healthy value is inventory minus capital at risk.", SANS, 13, BODY, False)])
footer(s, 10)

# ---------- 11 Roadmap ----------
s = prs.slides.add_slide(BLANK)
bg(s)
kicker(s, 11, "ROADMAP")
title_block(s, "What is built, and what waits on real use.")
cols = [
    ("Now", TEAL, [
        "Dead-stock and cover rules",
        "Reorder quantity from lead time",
        "Rescue prices at 10, 15, 20%",
        "Festival lift inside 14 days",
        "Voice that can record a sale",
        "Ledger of sales and purchases",
    ]),
    ("Next", GOLD, [
        "One database for a second device",
        "Accounts that keep a password",
        "Refit the risk score after weeks of sales",
        "A fuller purchase-order life cycle",
        "Invoice reading on more bill layouts",
    ]),
    ("Later", NAVY, [
        "Bundle prices beside discounts",
        "Supplier portals",
        "More than one shop on one book",
        "The proposed React and FastAPI stack, if the product leaves this browser",
    ]),
]
for i, (head, col, items) in enumerate(cols):
    x = 0.55 + i * 4.2
    rrect(s, x, 1.6, 3.95, 5.15, CARD, GOLD, 1.25)
    rect(s, x, 1.6, 3.95, 0.62, col)
    text(s, x + 0.2, 1.72, 3.55, 0.4, [(head, SERIF, 20, WHITE if col != GOLD else NAVY, False)])
    for j, item in enumerate(items):
        text(s, x + 0.25, 2.45 + j * 0.62, 3.45, 0.55, [(item, SANS, 14, NAVY, False)])
footer(s, 11)

# ---------- 12 Close ----------
s = prs.slides.add_slide(BLANK)
bg(s, NAVY)
rect(s, 0, 0, W, 0.08, GOLD)
rect(s, 0, 7.42, W, 0.08, GOLD)
mark(s, 0.7, 1.15, 0.7)
text(s, 1.6, 1.25, 8, 0.55, [("MERCURA AI", SERIF, 28, IVORY, False)], anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.7, 2.3, 12, 1.1, [("Better Inventory.", SERIF, 48, IVORY, False)])
text(s, 0.7, 3.2, 12, 1.0, [("Brighter Futures.", SERIF, 48, GOLD, False)])
rect(s, 0.72, 4.45, 2.0, 0.015, GOLD)
text(s, 0.7, 4.65, 11, 0.4, [("Driven by AI and controlled by you.", SANS, 18, TEAL, True)])
text(s, 0.7, 5.3, 11, 0.3, [("TEAM STACK OVERFLOW", SANS, 12, GOLD, True)])
text(s, 0.7, 5.65, 12, 0.35, [("Anusha K    ·    Abhishek Anil Kumar    ·    Aryan Singh    ·    Adarsh B Acharya", SANS, 16, IVORY, False)])
text(s, 0.7, 6.2, 12, 0.3, [("github.com/aryansingh7877/mercula-Ai", SANS, 14, GOLD, False)])
text(s, 0.7, 6.6, 12, 0.3, [("T.01 FinTech & Commerce    ·    Manipal Academy of Higher Education", SANS, 13, IVORY, False)])

out = "/Users/anushak/Documents/Projects/Hackathon/Mercura-App-main/Mercura-AI-Pitch.pptx"
prs.save(out)
print(out, len(prs.slides))
